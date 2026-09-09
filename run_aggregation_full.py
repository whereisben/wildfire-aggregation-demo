#!/usr/bin/env python3
"""
Wildfire 1km Exposure Aggregation - Geospatial Pipeline
============================================================================
Drop the property workbook in this folder and run:  python3 run_aggregation_full.py
Produces Wildfire_1km_Aggregation_FULL.xlsx with:
  - Aggregation Results (TIV): 6 filter versions + Worst-Case + equal-count quintiles
    (TIV & TIVxp(L)) + zone diff + optimal WC circle centre. Zones use NEW equal-width limits.
  - Aggregation Results (TIVxp(L)): single all/all optimisation maximising TIV*p(L).
  - Circle Summary (TIVxpL): canonical circles with member stats + expected loss.
  - WC Zone 5 Circles / WC Zone 5 Properties: canonical Zone-5 extracts.
  - Version Key & Heat Map.
Pivot summaries are computed value tables (native Excel PivotTables can't be authored by code).
If a hotspots globe HTML (e.g. *globe*.html) is in the folder, its embedded DATA is regenerated in place on every run.
Edit CONFIG only if the new file's tab / column names differ.
"""
import os, glob, math, json
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as GL

# ===== CONFIG =====
FOLDER=os.path.dirname(os.path.abspath(__file__))
DATASET_GLOB="*Synthetic*.xlsx"
SUMMARY="Summary"
OUTPUT="Wildfire_1km_Aggregation_FULL.xlsx"
GLOBE_GLOB=("*hotspots*globe*.html","*globe*.html","*hotspots*.html")  # regenerated in place if present
COLS=dict(ID="ID",LAT="Latitude",LNG="Longitude",SCORE="GRe WF Score",TIV="TIV",
          PF="p(f)",PLF="PLF",PL="p(L) final",TIVPL="TIV*p(L)",CITY="City",RETOOL="Retool Property ID")
VERSIONS={1:(None,None),4:(0.5,None),5:(0.5,0.5),7:(1.5,None),8:(1.5,0.5),9:(1.5,1.5)}
ORDER=[1,4,5,7,8,9]
LIMITS={1:150e6,4:150e6,5:100e6,7:150e6,8:100e6,9:50e6}   # NEW equal-width; edit to change
EARTH_RADIUS_KM=6371.0088; INNER=1.0; OUTER=2.0; CLIM=1.0; EPS=0.01
def hav(a,b,c,d):
    p1,p2=math.radians(a),math.radians(c); dp=math.radians(c-a); dl=math.radians(d-b)
    x=math.sin(dp/2)**2+math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return 2*EARTH_RADIUS_KM*math.atan2(math.sqrt(x),math.sqrt(1-x))
def brg(a,b,c,d):
    p1,p2=math.radians(a),math.radians(c); dl=math.radians(d-b)
    y=math.sin(dl)*math.cos(p2); x=math.cos(p1)*math.sin(p2)-math.sin(p1)*math.cos(p2)*math.cos(dl)
    return (math.degrees(math.atan2(y,x))+360)%360
def dest(a,b,bd,dk):
    ad=dk/EARTH_RADIUS_KM; bb=math.radians(bd); p1=math.radians(a); l1=math.radians(b)
    p2=math.asin(math.sin(p1)*math.cos(ad)+math.cos(p1)*math.sin(ad)*math.cos(bb))
    l2=l1+math.atan2(math.sin(bb)*math.sin(ad)*math.cos(p1),math.cos(ad)-math.sin(p1)*math.sin(p2))
    return {"lat":math.degrees(p2),"lng":((math.degrees(l2)+540)%360)-180}
def ci(p1,p2,r=INNER):
    d=hav(p1["lat"],p1["lng"],p2["lat"],p2["lng"])
    if d>2*r+EPS or d<EPS: return []
    if abs(d-2*r)<EPS:
        b=brg(p1["lat"],p1["lng"],p2["lat"],p2["lng"]); return [dest(p1["lat"],p1["lng"],b,r)]
    hc=math.sqrt(max(0,r*r-(d/2)**2)); b12=brg(p1["lat"],p1["lng"],p2["lat"],p2["lng"])
    mp=dest(p1["lat"],p1["lng"],b12,d/2)
    return [dest(mp["lat"],mp["lng"],(b12+90)%360,hc),dest(mp["lat"],mp["lng"],(b12-90+360)%360,hc)]
def optimize(home,neighbors,vk):
    pts=[{"lat":home["lat"],"lng":home["lng"]}]+[{"lat":n["lat"],"lng":n["lng"]} for n in neighbors]
    cand=list(pts)
    for i in range(len(pts)):
        for j in range(i+1,len(pts)): cand.extend(ci(pts[i],pts[j]))
    ded=[]; seen=set()
    for c in cand:
        k=(round(c["lat"],12),round(c["lng"],12))
        if k not in seen: seen.add(k); ded.append(c)
    def inside(c): return [n for n in neighbors if hav(n["lat"],n["lng"],c["lat"],c["lng"])<=INNER+EPS]
    bc={"lat":home["lat"],"lng":home["lng"]}; bn=inside(bc); bv=sum(n.get(vk,0) or 0 for n in bn)
    for c in ded:
        if hav(home["lat"],home["lng"],c["lat"],c["lng"])>CLIM+EPS: continue
        ins=inside(c); v=sum(n.get(vk,0) or 0 for n in ins)
        if v>bv: bc=c; bn=ins; bv=v
    hv=home.get(vk,0) or 0
    return {"ids":[home["id"]]+[n["id"] for n in bn],"neighbor_ids":[n["id"] for n in bn],
            "max_value":bv+hv,"center":bc}


def zone_ew(v,limit):
    if v<=0: return 0
    if v>=limit: return 5
    return min(int(v//(limit/4.0))+1,4)

def quintile(values):
    arr=np.array(values,dtype=float); N=len(arr)
    s=np.sort(arr); cnt=np.searchsorted(s,arr,side='right')
    return np.maximum(np.minimum(5,np.ceil(cnt/N*5)).astype(int),1)

def ids_str(ids,sep): return sep.join(str(x) for x in ids)
def num_or0(x): return 0.0 if pd.isna(x) else float(x)
def canon_key(s):
    parts=[p for p in str(s).replace(' ','').split(',') if p!='']
    return ','.join(str(x) for x in sorted(int(p) for p in parts)) if parts else ''
def canon_flags(keys):
    seen=set(); out=[]
    for k in keys:
        if k=='' or k in seen: out.append(0)
        else: seen.add(k); out.append(1)
    return out

def find_dataset():
    # 1) preferred name patterns (synthetic demo portfolio first)
    for pat in ("*Synthetic*.xlsx", "*Portfolio*.xlsx", "*Inforce*Risk*.xlsx"):
        hits=[h for h in glob.glob(os.path.join(FOLDER,pat)) if not os.path.basename(h).startswith("~")]
        if hits: return sorted(hits)[0]
    # 2) fallback: any .xlsx (excluding our own outputs) that has the SUMMARY tab with required cols
    skip={OUTPUT,"Wildfire_1km_Aggregation_Results.xlsx"}
    need={COLS["ID"],COLS["LAT"],COLS["LNG"],COLS["SCORE"],COLS["TIV"]}
    for p in sorted(glob.glob(os.path.join(FOLDER,"*.xlsx"))):
        if os.path.basename(p) in skip or os.path.basename(p).startswith("~$"): continue
        try:
            xl=pd.ExcelFile(p)
            if SUMMARY in xl.sheet_names:
                cols=set(pd.read_excel(p,sheet_name=SUMMARY,nrows=0).columns)
                if need.issubset(cols): return p
        except Exception:
            continue
    raise FileNotFoundError(
        f"No dataset found in {FOLDER}. Expected an .xlsx matching {DATASET_GLOB!r}, "
        f"or any .xlsx with a {SUMMARY!r} tab containing columns {sorted(need)}. "
        f"If the tab/columns are named differently, edit the CONFIG block (SUMMARY / COLS).")

def main():
    f=find_dataset(); print("Data:",os.path.basename(f))
    df=pd.read_excel(f,sheet_name=SUMMARY); df[COLS["ID"]]=df[COLS["ID"]].astype(int)
    _tp=pd.to_numeric(df[COLS["TIVPL"]],errors="coerce"); df=df[_tp.notna()&(_tp!=0)].reset_index(drop=True)  # only rows with nonzero TIV*p(L)
    base=list(df.columns)
    rows=[]
    for _,r in df.iterrows():
        rows.append(dict(id=int(r[COLS["ID"]]),lat=float(r[COLS["LAT"]]),lng=float(r[COLS["LNG"]]),
            tiv=num_or0(r[COLS["TIV"]]),score=num_or0(r[COLS["SCORE"]]),pf=num_or0(r[COLS["PF"]]),
            plf=num_or0(r[COLS["PLF"]]),pl=num_or0(r[COLS["PL"]]),tivpl=num_or0(r[COLS["TIVPL"]]),
            city=r[COLS["CITY"]],retool=r[COLS["RETOOL"]]))
    n=len(rows); byid={r["id"]:r for r in rows}
    lat=np.array([r["lat"] for r in rows]); lng=np.array([r["lng"] for r in rows])
    nidx=[]
    for i in range(n):
        m=(np.abs(lat-lat[i])<0.05)&(np.abs(lng-lng[i])<0.05)
        nidx.append([int(j) for j in np.where(m)[0] if j!=i and hav(lat[i],lng[i],lat[j],lng[j])<=OUTER+EPS])
    passes=lambda s,t: True if t is None else s>t
    print("Rows:",n)

    ver={v:[] for v in ORDER}
    for v in ORDER:
        rt,it=VERSIONS[v]
        for i,h in enumerate(rows):
            if not passes(h["score"],rt): ver[v].append(dict(ids='',val=0.0,zone=0,clat='',clng='')); continue
            nb=[dict(id=rows[j]["id"],lat=rows[j]["lat"],lng=rows[j]["lng"],tiv=rows[j]["tiv"])
                for j in nidx[i] if passes(rows[j]["score"],it)]
            res=optimize(dict(id=h["id"],lat=h["lat"],lng=h["lng"],tiv=h["tiv"]),nb,"tiv")
            ver[v].append(dict(ids=ids_str(res["ids"],','),val=res["max_value"],
                zone=zone_ew(res["max_value"],LIMITS[v]),clat=res["center"]["lat"],clng=res["center"]["lng"]))
    wc=[]
    for i in range(n):
        best=None
        for v in ORDER:
            dd=ver[v][i]; key=(dd["zone"],dd["val"],-v)
            if best is None or key>best[0]: best=(key,v)
        v=best[1]; dd=ver[v][i]
        wc.append(dict(zone=dd["zone"],version=f"V{v}",val=dd["val"],ids=dd["ids"],clat=dd["clat"],clng=dd["clng"]))
    txpl=[]
    for i,h in enumerate(rows):
        nb=[dict(id=rows[j]["id"],lat=rows[j]["lat"],lng=rows[j]["lng"],tivpl=rows[j]["tivpl"]) for j in nidx[i]]
        res=optimize(dict(id=h["id"],lat=h["lat"],lng=h["lng"],tivpl=h["tivpl"]),nb,"tivpl")
        txpl.append(dict(ids=ids_str(res["ids"],', '),nbr=ids_str(res["neighbor_ids"],', '),
            val=res["max_value"],count=len(res["ids"]),clat=res["center"]["lat"],clng=res["center"]["lng"]))
    tiv_q=list(map(int,quintile([w["val"] for w in wc])))
    txpl_q=list(map(int,quintile([t["val"] for t in txpl])))
    wc_canon=canon_flags([canon_key(w["ids"]) for w in wc])
    txpl_canon=canon_flags([canon_key(t["ids"]) for t in txpl])

    build_workbook(df,base,rows,byid,ver,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon)
    print("Saved:",OUTPUT)
    update_globe(FOLDER, build_globe_data(rows,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon))
    write_pivots(FOLDER, rows, byid, wc, txpl, tiv_q, txpl_q, wc_canon, txpl_canon)


def build_workbook(df,base,rows,byid,ver,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon):
    n=len(rows)
    VDEF={1:("all","all"),4:("gt 0.5","all"),5:("gt 0.5","gt 0.5"),7:("gt 1.5","all"),8:("gt 1.5","gt 0.5"),9:("gt 1.5","gt 1.5")}
    def members(ids): return [byid[int(p)] for p in str(ids).replace(' ','').split(',') if p!='']
    A=lambda **k:Font(name="Arial",**k)
    HF=PatternFill("solid",fgColor="1F3864");VF=PatternFill("solid",fgColor="2E5496");SF=PatternFill("solid",fgColor="595959")
    WF=PatternFill("solid",fgColor="7B3F00");QF=PatternFill("solid",fgColor="2F6B4F");XF=PatternFill("solid",fgColor="7030A0")
    thin=Side(style="thin",color="D9D9D9");bd=Border(thin,thin,thin,thin)
    def hdr(ws,r,c,val,fill=HF):
        x=ws.cell(r,c,val);x.font=A(bold=True,color="FFFFFF");x.fill=fill;x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    def put(ws,r,c,val,num=None,ctr=False):
        x=ws.cell(r,c,val);x.font=A(size=10);x.border=bd;x.alignment=Alignment(horizontal="center" if ctr else "left",vertical="center")
        if num:x.number_format=num
    def bnum(name):
        return '#,##0' if name==COLS["TIV"] else ('#,##0.00' if name==COLS["TIVPL"] else ('0.000000' if name in(COLS["LAT"],COLS["LNG"],COLS["PF"],COLS["PLF"],COLS["PL"]) else ('0.00' if name==COLS["SCORE"] else ('0' if name in("Zip",COLS["ID"],"Policy Number") else None))))
    
    wb=Workbook()
    # ---- SHEET 0: Cover & Disclaimer ----
    ws0=wb.active; ws0.title="Disclaimer & Overview"
    ws0.cell(1,1,"Wildfire 1km Exposure Aggregation Model").font=A(size=15,bold=True,color="1F3864")
    ws0.cell(2,1,"DEMONSTRATION PORTFOLIO · 100% SYNTHETIC DATASET").font=A(size=11,bold=True,color="C00000")
    desc = [
        ("Disclaimer:", "All properties, coordinates, policy numbers, and insured values in this model are 100% synthetically generated."),
        ("Privacy:", "Contains ZERO company proprietary data or real policyholder information. Built specifically for code portfolio showcase."),
        ("Methodology:", "Spatial aggregation finding the optimal 1km circle (free center) maximizing exposure across 6 wildfire filter versions."),
        ("Heat Map:", "Equal-width quartiles [0, Limit] with Zone 5 >= Limit (Limits: 150M for all, 100M for >0.5, 50M for >1.5)."),
        ("Quintiles:", "Equal-count quintiles (Rank method) across Worst-Case TIV and TIV*p(L)."),
        ("Outputs:", "8 analysis sheets including Worst-Case rollups, limit splits, and canonical circle extracts."),
        ("Simulated Properties:", f"{len(rows)} across {df[COLS['CITY']].nunique()} California cities"),
        ("Portfolio TIV:", f"${df[COLS['TIV']].sum():,.0f}"),
    ]
    for r_idx, (k, v) in enumerate(desc, start=4):
        ws0.cell(r_idx, 1, k).font = A(bold=True, size=10, color="1F3864")
        ws0.cell(r_idx, 2, v).font = A(size=10)
    ws0.column_dimensions["A"].width = 24; ws0.column_dimensions["B"].width = 90
    
    # ---- SHEET 1: Aggregation Results (TIV) ----
    ws=wb.create_sheet("Aggregation Results (TIV)")
    cols=list(base)
    for v in ORDER: cols+=[f"V{v} Circle IDs",f"V{v} Max Total TIV",f"V{v} Zone"]
    cols+=["WC Zone","WC Version","WC Max Total TIV","WC Circle IDs","TIV Quintile Zone","TIVxp(L) Quintile Zone","Zone Diff (TIVxp(L) - TIV)","Optimal Center Latitude","Optimal Center Longitude"]
    c=1; ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(base)); hdr(ws,1,1,"Original Summary columns (source data)",SF)
    for name in base: hdr(ws,2,c,name); c+=1
    for v in ORDER:
        ws.merge_cells(start_row=1,start_column=c,end_row=1,end_column=c+2); rf,inf=VDEF[v]
        hdr(ws,1,c,f"Version {v} (row: {rf} | incl: {inf} | limit {int(LIMITS[v]/1e6)}m)",VF)
        for k,nm in enumerate(["Circle IDs","Max Total TIV","Zone"]): hdr(ws,2,c+k,nm)
        c+=3
    ws.merge_cells(start_row=1,start_column=c,end_row=1,end_column=c+3); hdr(ws,1,c,"Worst-Case (max zone across 6 versions)",WF)
    for k,nm in enumerate(["WC Zone","WC Version","WC Max Total TIV","WC Circle IDs"]): hdr(ws,2,c+k,nm)
    c+=4; ws.merge_cells(start_row=1,start_column=c,end_row=1,end_column=c+2); hdr(ws,1,c,"Quintiles (equal-count) & diff",QF)
    for k,nm in enumerate(["TIV Quintile Zone","TIVxp(L) Quintile Zone","Zone Diff (TIVxp(L) - TIV)"]): hdr(ws,2,c+k,nm)
    c+=3; ws.merge_cells(start_row=1,start_column=c,end_row=1,end_column=c+1); hdr(ws,1,c,"Optimal WC Circle Center",XF)
    for k,nm in enumerate(["Latitude","Longitude"]): hdr(ws,2,c+k,nm)
    for i,(_,r) in enumerate(df.iterrows()):
        ri=i+3; c=1
        for name in base: put(ws,ri,c,r[name],bnum(name)); c+=1
        for v in ORDER:
            dd=ver[v][i]; put(ws,ri,c,dd["ids"]); put(ws,ri,c+1,dd["val"],'#,##0'); put(ws,ri,c+2,dd["zone"],ctr=True); c+=3
        w=wc[i]; put(ws,ri,c,w["zone"],ctr=True); put(ws,ri,c+1,w["version"],ctr=True); put(ws,ri,c+2,w["val"],'#,##0'); put(ws,ri,c+3,w["ids"]); c+=4
        put(ws,ri,c,tiv_q[i],ctr=True); put(ws,ri,c+1,txpl_q[i],ctr=True); put(ws,ri,c+2,txpl_q[i]-tiv_q[i],ctr=True); c+=3
        put(ws,ri,c,w["clat"],'0.00000000' if w["clat"]!='' else None); put(ws,ri,c+1,w["clng"],'0.00000000' if w["clng"]!='' else None)
    for name in cols:
        L=GL(cols.index(name)+1)
        if name.endswith("Circle IDs"): ws.column_dimensions[L].width=24
        elif "Max Total TIV" in name: ws.column_dimensions[L].width=14
        elif name.endswith("Zone") or "Quintile" in name or "Diff" in name: ws.column_dimensions[L].width=9
        elif name==COLS["RETOOL"]: ws.column_dimensions[L].width=30
        elif name in(COLS["LAT"],COLS["LNG"]): ws.column_dimensions[L].width=13
        else: ws.column_dimensions[L].width=12
    ws.freeze_panes="B3"; ws.row_dimensions[1].height=30; ws.row_dimensions[2].height=28

    ws2=wb.create_sheet("Aggregation Results (TIVxp(L))")
    h2=list(base)+["Max Total TIV*p(L)","IDs in Optimal 1km Circle","Count of IDs","Neighbor IDs (excl. home)","Optimal Center Latitude","Optimal Center Longitude"]
    for j,nm in enumerate(h2,start=1): hdr(ws2,1,j,nm)
    for i,(_,r) in enumerate(df.iterrows()):
        ri=i+2; c=1
        for name in base: put(ws2,ri,c,r[name],bnum(name)); c+=1
        t=txpl[i]; put(ws2,ri,c,t["val"],'"$"#,##0.00'); put(ws2,ri,c+1,t["ids"]); put(ws2,ri,c+2,t["count"],ctr=True); put(ws2,ri,c+3,t["nbr"]); put(ws2,ri,c+4,t["clat"],'0.00000000'); put(ws2,ri,c+5,t["clng"],'0.00000000')
    for j,nm in enumerate(h2,start=1):
        L=GL(j)
        ws2.column_dimensions[L].width=24 if "IDs" in nm else (30 if nm==COLS["RETOOL"] else 13)
    ws2.freeze_panes="A2"; ws2.row_dimensions[1].height=28

    ws3=wb.create_sheet("Circle Summary (TIVxpL)")
    ws3.cell(1,1,"Canonical circles from the TIV*p(L) optimisation (one row per unique circle)").font=A(bold=True,size=11)
    h3=["Home ID","City","# of Properties","Avg p(f)","Avg PLF","Avg p(L)","Avg GRE Score","Sum of TIV","Risk Adjusted TIV (Sum TIV*p(L))"]
    for j,nm in enumerate(h3,start=1): hdr(ws3,2,j,nm)
    rr=3
    for i,r in enumerate(rows):
        if not txpl_canon[i]: continue
        mem=members(txpl[i]["ids"])
        put(ws3,rr,1,r["id"],ctr=True); put(ws3,rr,2,r["city"]); put(ws3,rr,3,len(mem),ctr=True)
        put(ws3,rr,4,float(np.mean([m["pf"] for m in mem])),'0.000000'); put(ws3,rr,5,float(np.mean([m["plf"] for m in mem])),'0.000000')
        put(ws3,rr,6,float(np.mean([m["pl"] for m in mem])),'0.000000'); put(ws3,rr,7,float(np.mean([m["score"] for m in mem])),'0.00')
        put(ws3,rr,8,float(np.sum([m["tiv"] for m in mem])),'#,##0'); put(ws3,rr,9,float(np.sum([m["tivpl"] for m in mem])),'"$"#,##0.00'); rr+=1
    for j in range(1,10): ws3.column_dimensions[GL(j)].width=10 if j==1 else 16
    ws3.freeze_panes="A3"

    ws4=wb.create_sheet("WC Zone 5 Circles")
    z5=[]
    for i,r in enumerate(rows):
        if wc[i]["zone"]==5 and wc_canon[i]:
            mem=members(wc[i]["ids"])
            z5.append((r["id"],r["city"],5,len(mem),float(np.mean([m["score"] for m in mem])),float(np.sum([m["tiv"] for m in mem])),wc[i]["ids"]))
    z5.sort(key=lambda x:-x[5])
    ws4.cell(1,1,f"Canonical Worst-Case circles in Zone 5 ({len(z5)} circles), sorted by Sum of TIV").font=A(bold=True,size=11)
    for j,nm in enumerate(["Home ID","City","WC Zone","# of Properties","Avg GRE Score","Sum of TIV","WC Circle IDs"],start=1): hdr(ws4,2,j,nm)
    for k,row in enumerate(z5):
        rr=k+3; put(ws4,rr,1,row[0],ctr=True); put(ws4,rr,2,row[1]); put(ws4,rr,3,row[2],ctr=True); put(ws4,rr,4,row[3],ctr=True)
        put(ws4,rr,5,row[4],'0.00'); put(ws4,rr,6,row[5],'#,##0'); put(ws4,rr,7,row[6])
    for j,w_ in enumerate([10,16,8,12,12,14,24],start=1): ws4.column_dimensions[GL(j)].width=w_
    ws4.freeze_panes="A3"

    ws5=wb.create_sheet("WC Zone 5 Properties")
    for j,nm in enumerate(["Circle Home ID","Property ID","City","Retool Property ID"],start=1): hdr(ws5,1,j,nm)
    rr=2
    for row in z5:
        for m in members(row[6]): put(ws5,rr,1,row[0],ctr=True); put(ws5,rr,2,m["id"],ctr=True); put(ws5,rr,3,m["city"]); put(ws5,rr,4,m["retool"]); rr+=1
    for j,w_ in enumerate([12,10,16,34],start=1): ws5.column_dimensions[GL(j)].width=w_
    ws5.freeze_panes="A2"

    ws6=wb.create_sheet("Version Key & Heat Map")
    ws6.cell(1,1,"Version definitions & NEW equal-width heat map").font=A(bold=True,size=12)
    for j,x in enumerate(["Version","Row filter","Include filter","Zone-5 limit","Z1","Z2","Z3","Z4","Z5"],start=1): hdr(ws6,2,j,x)
    mm=lambda x:f"{x/1e6:g}m"
    for i,v in enumerate(ORDER,start=3):
        rf,inf=VDEF[v]; L=LIMITS[v]; q=L/4
        for j,x in enumerate([v,rf,inf,mm(L),f"0-{mm(q)}",f"{mm(q)}-{mm(2*q)}",f"{mm(2*q)}-{mm(3*q)}",f"{mm(3*q)}-{mm(L)}",f">= {mm(L)}"],start=1): put(ws6,i,j,x,ctr=True)
    for j,w_ in enumerate([8,10,12,11,12,12,13,14,10],start=1): ws6.column_dimensions[GL(j)].width=w_

    ws7=wb.create_sheet("WC Zone 5 by Limit")
    ws7.cell(1,1,"Worst-Case Zone-5 homes, one table per Max TIV limit (desc by Max Total TIV)").font=A(bold=True,size=12)
    z5rows=[]
    for i,r in enumerate(rows):
        w=wc[i]
        if w["zone"]==5:
            z5rows.append((LIMITS[int(w["version"][1:])],r,w))
    rr=3
    for lim in (150e6,100e6,50e6):
        grp=sorted([(r,w) for L,r,w in z5rows if L==lim],key=lambda x:-x[1]["val"])
        hdr(ws7,rr,1,f"{mm(lim)} limit  ({len(grp)} homes)",WF)
        for k in range(2,7): hdr(ws7,rr,k,"",WF)
        rr+=1
        for j,nm in enumerate(["Home ID","City","WC Version","Max Total TIV","# of Properties","WC Circle IDs"],start=1): hdr(ws7,rr,j,nm)
        rr+=1
        for r,w in grp:
            put(ws7,rr,1,r["id"],ctr=True); put(ws7,rr,2,r["city"]); put(ws7,rr,3,w["version"],ctr=True)
            put(ws7,rr,4,w["val"],'#,##0'); put(ws7,rr,5,len(members(w["ids"])),ctr=True); put(ws7,rr,6,w["ids"]); rr+=1
        rr+=1
    for j,w_ in enumerate([10,16,11,15,13,24],start=1): ws7.column_dimensions[GL(j)].width=w_

    ws8=wb.create_sheet("WC Top Quintile by Limit")
    ws8.cell(1,1,"Top-quintile homes (top 20% highest Max Total TIV) per Max TIV limit, desc").font=A(bold=True,size=12)
    rr=3
    for lim in (150e6,100e6,50e6):
        grp=[(r,wc[i]) for i,r in enumerate(rows) if LIMITS[int(wc[i]["version"][1:])]==lim]
        q=quintile([w["val"] for _,w in grp])
        top=sorted([grp[k] for k in range(len(grp)) if q[k]==5],key=lambda x:-x[1]["val"])
        hdr(ws8,rr,1,f"{mm(lim)} limit  (top 20% = {len(top)} of {len(grp)} homes)",QF)
        for k in range(2,7): hdr(ws8,rr,k,"",QF)
        rr+=1
        for j,nm in enumerate(["Home ID","City","WC Version","Max Total TIV","# of Properties","WC Circle IDs"],start=1): hdr(ws8,rr,j,nm)
        rr+=1
        for r,w in top:
            put(ws8,rr,1,r["id"],ctr=True); put(ws8,rr,2,r["city"]); put(ws8,rr,3,w["version"],ctr=True)
            put(ws8,rr,4,w["val"],'#,##0'); put(ws8,rr,5,len(members(w["ids"])),ctr=True); put(ws8,rr,6,w["ids"]); rr+=1
        rr+=1
    for j,w_ in enumerate([10,16,11,15,13,24],start=1): ws8.column_dimensions[GL(j)].width=w_

    wb.save(os.path.join(FOLDER,OUTPUT))


def _memids(s):
    return [int(x) for x in str(s).replace(" ","").split(",") if x!=""]

def build_globe_data(rows,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon):
    homes={str(r["id"]):[round(r["lat"],6),round(r["lng"],6),r["tiv"],int(txpl_q[i]-tiv_q[i]),round(r["tivpl"],2)] for i,r in enumerate(rows)}
    pl_by_id={r["id"]:r["tivpl"] for r in rows}
    tiv_by_id={r["id"]:r["tiv"] for r in rows}
    tiv=[]
    for i,r in enumerate(rows):
        if not wc_canon[i]: continue
        w=wc[i]; mem=_memids(w["ids"]) or [r["id"]]
        tiv.append({"lat":round(w["clat"],6) if w["clat"]!="" else round(r["lat"],6),
                    "lon":round(w["clng"],6) if w["clng"]!="" else round(r["lng"],6),
                    "v":round(w["val"],2),"el":round(sum(pl_by_id[m] for m in mem),2),
                    "qz":int(tiv_q[i]),"n":len(mem),"city":r["city"],"ids":mem,"wz":int(w["zone"])})
    tivpl=[]
    for i,r in enumerate(rows):
        if not txpl_canon[i]: continue
        t=txpl[i]; mem=_memids(t["ids"]) or [r["id"]]
        tivpl.append({"lat":round(t["clat"],6),"lon":round(t["clng"],6),"v":round(t["val"],2),
                      "tv":round(sum(tiv_by_id[m] for m in mem),2),
                      "qz":int(txpl_q[i]),"n":t["count"],"city":r["city"],"ids":mem})

    def _nms_circles(circles, dist_thresh=1.2, iou_thresh=0.30):
        # Non-maximum suppression: keep highest-exposure circles and suppress redundant overlaps
        kept = []
        for c in sorted(circles, key=lambda x: (x["v"], x["n"]), reverse=True):
            c_mem = set(c["ids"])
            suppress = False
            for k in kept:
                d = hav(c["lat"], c["lon"], k["lat"], k["lon"])
                if d < dist_thresh:
                    k_mem = set(k["ids"])
                    inter = len(c_mem & k_mem)
                    if inter > 0:
                        overlap = inter / max(1, min(len(c_mem), len(k_mem)))
                        if overlap >= iou_thresh or d < 0.8:
                            suppress = True
                            break
            if not suppress:
                kept.append(c)
        return kept

    return {"homes":homes,"tiv":_nms_circles(tiv),"tivpl":_nms_circles(tivpl)}

def update_globe(folder,data):
    path=None
    for pat in GLOBE_GLOB:
        hits=glob.glob(os.path.join(folder,pat))
        if hits: path=sorted(hits)[0]; break
    if not path: print("  globe template not found - skipping globe update"); return
    s=open(path,encoding="utf-8").read(); m=s.find("const DATA=")
    if m<0: print("  no 'const DATA=' marker in globe - skipping"); return
    start=m+len("const DATA="); depth=0; end=None
    for k in range(start,len(s)):
        if s[k]=="{": depth+=1
        elif s[k]=="}":
            depth-=1
            if depth==0: end=k+1; break
    open(path,"w",encoding="utf-8").write(s[:start]+json.dumps(data,separators=(",",":"))+s[end:])
    print("  globe updated: %s (tiv=%d, tivpl=%d, homes=%d)"%(os.path.basename(path),len(data["tiv"]),len(data["tivpl"]),len(data["homes"])))



def pivots_html(recs):
    data_js=json.dumps(recs,separators=(',',':'))
    tpl = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Wildfire 1km Aggregation Pivots (Demo Portfolio)</title>
<style>
:root{--bg:#0f1115;--panel:#171a21;--line:#262b34;--fg:#e7eaf0;--mut:#9aa3b2;--acc:#FF6500;}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.45 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif}
header{padding:18px 22px;border-bottom:1px solid var(--line);display:flex;align-items:baseline;gap:14px;flex-wrap:wrap}
h1{font-size:18px;margin:0;font-weight:650;letter-spacing:.2px}
.sub{color:var(--mut);font-size:12.5px}
.wrap{padding:18px 22px;max-width:1300px}
.controls{display:flex;gap:18px;flex-wrap:wrap;align-items:flex-end;margin-bottom:16px}
.ctl label{display:block;font-size:11px;text-transform:uppercase;letter-spacing:.6px;color:var(--mut);margin-bottom:5px}
select{background:var(--panel);color:var(--fg);border:1px solid var(--line);border-radius:7px;padding:7px 10px;font-size:13px;min-width:150px}
.cards{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:16px}
.card{background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:12px 16px;min-width:150px}
.card .k{font-size:11px;text-transform:uppercase;letter-spacing:.6px;color:var(--mut)}
.card .v{font-size:22px;font-weight:650;margin-top:3px}
table{width:100%;border-collapse:collapse;background:var(--panel);border:1px solid var(--line);border-radius:10px;overflow:hidden}
th,td{padding:8px 11px;text-align:right;border-bottom:1px solid var(--line);white-space:nowrap}
th:first-child,td:first-child{text-align:left}
th{position:sticky;top:0;background:#1c2029;font-size:11px;text-transform:uppercase;letter-spacing:.5px;color:var(--mut);cursor:pointer;user-select:none}
th.sorted{color:var(--acc)}
tr:hover td{background:#1b1f27}
.z{display:inline-block;min-width:20px;text-align:center;border-radius:5px;padding:1px 7px;color:#fff;font-weight:600;font-size:12px}
.tblwrap{max-height:70vh;overflow:auto;border-radius:10px}
.ids{color:var(--mut);font-size:12px;max-width:280px;overflow:hidden;text-overflow:ellipsis}
.note{color:var(--mut);font-size:12px;margin-top:10px}
</style></head><body>
<header>
  <div>
    <h1>Wildfire Risk — 1 km Aggregation Pivots</h1>
    <span class="sub">Canonical circles (one row per unique circle) · click a column to sort</span>
  </div>
  <div style="background:#432b00;border:1px solid #d48806;color:#ffe58f;padding:6px 12px;border-radius:6px;font-size:12px;font-weight:600;display:inline-flex;align-items:center;gap:6px;">
    <span>⚠️</span> <span>DEMO PORTFOLIO · 100% SYNTHETIC DATASET</span>
  </div>
</header>
<div class="wrap">
  <div class="controls">
    <div class="ctl"><label>Dataset</label>
      <select id="dataset">
        <option value="wc">Worst-Case (TIV)</option>
        <option value="tivpl">Risk Adjusted TIV (TIV×p(L))</option>
      </select></div>
    <div class="ctl"><label>Zone field</label>
      <select id="zfield"></select></div>
    <div class="ctl"><label>Zone filter</label>
      <select id="zval">
        <option value="all">All zones</option>
        <option value="5" selected>Zone 5 only</option>
        <option value="4">Zone 4+</option>
      </select></div>
    <div class="ctl"><label>Group by</label>
      <select id="group">
        <option value="circle">Circle (each)</option>
        <option value="city">City</option>
        <option value="zone">Zone</option>
      </select></div>
  </div>
  <div class="cards" id="cards"></div>
  <div class="tblwrap"><table id="tbl"><thead></thead><tbody></tbody></table></div>
  <div class="note" id="note"></div>
</div>
<script>
const PIVOT_DATA=__DATA__;
const ZC={1:'#595959',2:'#8a5f43',3:'#bd6a2a',4:'#e0630f',5:'#FF6500'};
const S={dataset:'wc',zfield:'wz',zval:'5',group:'circle',sortKey:'sum_tiv',sortDir:-1};
const $=id=>document.getElementById(id);
const money=v=>v>=1e9?'$'+(v/1e9).toFixed(2)+'B':v>=1e6?'$'+(v/1e6).toFixed(1)+'M':v>=1e3?'$'+Math.round(v/1e3)+'K':'$'+Math.round(v);
const zchip=z=>`<span class="z" style="background:${ZC[z]||'#888'}">${z}</span>`;

function zfieldOpts(){
  const el=$('zfield'); el.innerHTML='';
  const opts = S.dataset==='wc' ? [['wz','WC Zone'],['qz','Quintile Zone']] : [['qz','Quintile Zone']];
  opts.forEach(([v,l])=>{const o=document.createElement('option');o.value=v;o.textContent=l;el.appendChild(o);});
  if(!opts.some(o=>o[0]===S.zfield)) S.zfield=opts[0][0];
  el.value=S.zfield;
}
function baseRecs(){
  let r=PIVOT_DATA[S.dataset].slice();
  if(S.zval!=='all'){const t=+S.zval; r=r.filter(x=>x[S.zfield]>=t && (S.zval==='5'?x[S.zfield]===5:true));}
  return r;
}
function wavg(rows,key){const num=rows.reduce((a,x)=>a+x[key]*x.n,0),den=rows.reduce((a,x)=>a+x.n,0);return den?num/den:0;}
function grouped(){
  const r=baseRecs();
  if(S.group==='circle') return {mode:'circle',rows:r};
  const keyf = S.group==='city' ? (x=>x.city) : (x=>x[S.zfield]);
  const m=new Map();
  r.forEach(x=>{const k=keyf(x); if(!m.has(k))m.set(k,[]); m.get(k).push(x);});
  const rows=[...m.entries()].map(([k,g])=>({
    key:k, circles:g.length, props:g.reduce((a,x)=>a+x.n,0),
    sum_tiv:g.reduce((a,x)=>a+x.sum_tiv,0), exp_loss:g.reduce((a,x)=>a+x.exp_loss,0),
    avg_gre:wavg(g,'avg_gre'), avg_pf:wavg(g,'avg_pf'), avg_plf:wavg(g,'avg_plf'), avg_pl:wavg(g,'avg_pl')
  }));
  rows.forEach(x=>x.avg_tiv=x.props?x.sum_tiv/x.props:0);
  return {mode:'group',rows};
}
const COLS={
  circle:[['home','Home ID',0],['city','City',0],['zoneCell','Zone',0],['n','Props',1],
    ['sum_tiv','Sum TIV',2],['avg_tiv','Avg TIV/home',2],['avg_gre','Avg GRE',3],
    ['avg_pf','Avg p(f)',4],['avg_plf','Avg PLF',4],['avg_pl','Avg p(L)',4],['exp_loss','Risk Adjusted TIV',2],['idsCell','Member IDs',0]],
  group:[['key','Group',0],['circles','Circles',1],['props','Props',1],['sum_tiv','Sum TIV',2],
    ['avg_tiv','Avg TIV/home',2],['avg_gre','Avg GRE',3],['avg_pf','Avg p(f)',4],
    ['avg_plf','Avg PLF',4],['avg_pl','Avg p(L)',4],['exp_loss','Risk Adjusted TIV',2]]
};
function fmt(v,t){return t===2?money(v):t===1?(''+v):t===3?v.toFixed(2):t===4?v.toFixed(5):v;}
function render(){
  zfieldOpts();
  const g=grouped(); const cols=COLS[g.mode];
  // default sort key valid?
  if(!cols.some(c=>c[0]===S.sortKey)) {S.sortKey=g.mode==='group'?'sum_tiv':'sum_tiv';}
  const sortable=new Set(cols.map(c=>c[0])).has(S.sortKey)?S.sortKey:'sum_tiv';
  g.rows.sort((a,b)=>{const x=a[sortable],y=b[sortable];return (x<y?-1:x>y?1:0)*S.sortDir;});
  // cards
  const base=baseRecs();
  const cir=base.length, props=base.reduce((a,x)=>a+x.n,0), tiv=base.reduce((a,x)=>a+x.sum_tiv,0), el=base.reduce((a,x)=>a+x.exp_loss,0);
  $('cards').innerHTML=[['Unique circles',cir],['Properties',props],['Sum of TIV',money(tiv)],['Risk Adjusted TIV',money(el)]]
    .map(([k,v])=>`<div class="card"><div class="k">${k}</div><div class="v">${v}</div></div>`).join('');
  // head
  $('tbl').tHead.innerHTML='<tr>'+cols.map(c=>`<th data-k="${c[0]}" class="${c[0]===sortable?'sorted':''}">${c[1]}${c[0]===sortable?(S.sortDir<0?' ▼':' ▲'):''}</th>`).join('')+'</tr>';
  // body
  const body=g.rows.map(row=>{
    const cells=cols.map(c=>{
      const k=c[0];
      if(k==='zoneCell'){return `<td>${zchip(row.wz??row[S.zfield])}${(S.dataset==='wc')?' <span style="color:var(--mut);font-size:11px">q'+row.qz+'</span>':''}</td>`;}
      if(k==='idsCell'){return `<td class="ids">${row.ids.join(', ')}</td>`;}
      if(k==='key'){return `<td>${S.group==='zone'?zchip(row.key):row.key}</td>`;}
      return `<td>${fmt(row[k],c[2])}</td>`;
    });
    return '<tr>'+cells.join('')+'</tr>';
  }).join('');
  $('tbl').tBodies[0].innerHTML=body;
  $('note').textContent=`${g.rows.length} ${g.mode==='circle'?'circles':'groups'} shown · dataset: ${S.dataset==='wc'?'Worst-Case TIV':'Risk Adjusted TIV (TIV×p(L))'} · filter: ${S.zfield.toUpperCase()} ${S.zval==='all'?'all':S.zval==='5'?'=5':'>=4'}`;
  document.querySelectorAll('th').forEach(th=>th.onclick=()=>{const k=th.dataset.k;if(S.sortKey===k)S.sortDir*=-1;else{S.sortKey=k;S.sortDir=-1;}render();});
}
['dataset','zfield','zval','group'].forEach(id=>$(id).addEventListener('change',e=>{
  S[id==='zfield'?'zfield':id==='zval'?'zval':id]=e.target.value; if(id==='dataset'){S.sortKey='sum_tiv';S.sortDir=-1;} render();
}));
render();
</script></body></html>"""
    return tpl.replace('__DATA__', data_js)


def _circle_records(rows,byid,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon):
    import numpy as _np
    def _mem(ids): return [byid[int(p)] for p in str(ids).replace(" ","").split(",") if p!=""]
    def _st(members):
        tiv=[m["tiv"] for m in members]
        return dict(n=len(members),sum_tiv=float(_np.sum(tiv)),avg_tiv=float(_np.mean(tiv)),
            avg_gre=float(_np.mean([m["score"] for m in members])),avg_pf=float(_np.mean([m["pf"] for m in members])),
            avg_plf=float(_np.mean([m["plf"] for m in members])),avg_pl=float(_np.mean([m["pl"] for m in members])),
            exp_loss=float(_np.sum([m["tivpl"] for m in members])))
    def _r(v): return round(v,4) if isinstance(v,float) else v
    wc_recs=[]
    for i,r in enumerate(rows):
        if not wc_canon[i]: continue
        m=_mem(wc[i]["ids"]) or [r]; s=_st(m)
        wc_recs.append(dict(home=r["id"],city=r["city"],wz=int(wc[i]["zone"]),qz=int(tiv_q[i]),
            val=round(wc[i]["val"],2),ids=[x["id"] for x in m],
            clat=float(wc[i]["clat"]) if wc[i]["clat"]!="" else r["lat"],
            clng=float(wc[i]["clng"]) if wc[i]["clng"]!="" else r["lng"],
            **{k:_r(v) for k,v in s.items()}))
    tp_recs=[]
    for i,r in enumerate(rows):
        if not txpl_canon[i]: continue
        m=_mem(txpl[i]["ids"]) or [r]; s=_st(m)
        tp_recs.append(dict(home=r["id"],city=r["city"],qz=int(txpl_q[i]),
            val=round(txpl[i]["val"],2),ids=[x["id"] for x in m],
            clat=float(txpl[i]["clat"]) if txpl[i]["clat"]!="" else r["lat"],
            clng=float(txpl[i]["clng"]) if txpl[i]["clng"]!="" else r["lng"],
            **{k:_r(v) for k,v in s.items()}))

    def _nms_recs(recs, dist_thresh=1.2, iou_thresh=0.30):
        kept = []
        for c in sorted(recs, key=lambda x: (x["val"], x["n"]), reverse=True):
            c_mem = set(c["ids"])
            suppress = False
            for k in kept:
                d = hav(c["clat"], c["clng"], k["clat"], k["clng"])
                if d < dist_thresh:
                    k_mem = set(k["ids"])
                    inter = len(c_mem & k_mem)
                    if inter > 0:
                        overlap = inter / max(1, min(len(c_mem), len(k_mem)))
                        if overlap >= iou_thresh or d < 0.8:
                            suppress = True
                            break
            if not suppress:
                kept.append(c)
        return kept

    return {"wc":_nms_recs(wc_recs),"tivpl":_nms_recs(tp_recs)}

def write_pivots(folder,rows,byid,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon,filename="aggregation_pivots.html"):
    recs=_circle_records(rows,byid,wc,txpl,tiv_q,txpl_q,wc_canon,txpl_canon)
    open(os.path.join(folder,filename),"w",encoding="utf-8").write(pivots_html(recs))
    print("  pivots dashboard written: %s (wc=%d, tivpl=%d)"%(filename,len(recs["wc"]),len(recs["tivpl"])))


if __name__=="__main__":
    main()
